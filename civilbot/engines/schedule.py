"""Column schedule: design many columns, size them, group identical designs into types."""
from __future__ import annotations
import math
import re
from ..core import Result, InputError, fmt
from . import rc_column


def parse(text: str):
    """Lines like 'C1 Pu=1500 Mu=80' or 'C1,1500,80' (Mu optional)."""
    cols = []
    for line in re.split(r"[;\n]+", text.strip()):
        line = line.strip()
        if not line:
            continue
        m = re.match(r"([A-Za-z]*\d+[A-Za-z]*)\s*[, ]\s*(.*)", line)
        if not m:
            continue
        name, rest = m.group(1), m.group(2)
        kv = {k.lower(): float(v) for k, v in re.findall(r"([A-Za-z]+)\s*=\s*([-\d.]+)", rest)}
        if not kv:
            nums = [float(x) for x in re.findall(r"[-\d.]+", rest)]
            if not nums:
                continue
            kv = {"pu": nums[0], "mu": nums[1] if len(nums) > 1 else 0.0}
        if "pu" not in kv:
            raise InputError(f"العمود {name}: ناقص Pu")
        cols.append((name.upper(), kv["pu"], kv.get("mu", 0.0), kv.get("n", 1)))
    if not cols:
        raise InputError("مثال: C1 Pu=1500 Mu=80 ; C2 Pu=2200 Mu=60")
    return cols


def design_all(cols, fc=28.0, fy=420.0, sizes=None, rho_max=0.03, cover=40.0, aspect=1.0, height=3.2) -> Result:
    sizes = sizes or list(range(300, 1001, 50))
    r = Result("schedule", "Column schedule (ACI 318-19)", inputs=dict(columns=len(cols), fc=f"{fmt(fc)} MPa", fy=f"{fmt(fy)} MPa",
                                                                       rule=f"smallest square/rect section with ρ ≤ {fmt(rho_max*100)} %"))
    designs = []
    for name, Pu, Mu, cnt in cols:
        best = None
        for b in sizes:
            h = math.ceil(b * aspect / 50) * 50
            try:
                d = rc_column.design(b, h, fc, fy, Pu, Mu, cover=cover)
            except InputError:
                continue
            if d.outputs["rho"] <= rho_max:
                best = d.outputs
                break
        if not best:
            raise InputError(f"العمود {name} يحتاج مقطع أكبر من {sizes[-1]} mm")
        designs.append((name, Pu, Mu, int(cnt), best))
    groups = {}
    for name, Pu, Mu, cnt, o in designs:
        key = (o["b"], o["h"], o["n"], o["db"])
        groups.setdefault(key, []).append((name, Pu, Mu, cnt, o))
    types = []
    for k, (key, items) in enumerate(sorted(groups.items(), key=lambda t: (t[0][0] * t[0][1], t[0][2] * t[0][3] ** 2)), 1):
        b, h, n, db = key
        o = items[0][4]
        tname = f"C{k}"
        count = sum(i[3] for i in items)
        conc = b * h / 1e6 * height * count
        steel = n * math.pi * db**2 / 4 / 1e6 * (height + 40 * db / 1000) * 7850 * count
        types.append(dict(type=tname, b=b, h=h, n=n, db=db, tie=o["tie"], s=o["s_tie"], members=[i[0] for i in items], count=count,
                          Pmax=max(i[1] for i in items), rho=o["rho"], conc=conc, steel=steel, o=o))
        r.add(f"Type {tname}: {fmt(b)}×{fmt(h)} — {n}Ø{db} — ties Ø{o['tie']}@{o['s_tie']}", "",
              f"columns: {', '.join(i[0] for i in items)} (×{count}) , Pu,max = {fmt(max(i[1] for i in items))} kN , ρ = {fmt(o['rho']*100, 2)} %",
              f"concrete {fmt(conc, 2)} m³ , steel ≈ {fmt(steel, 0)} kg")
    tot_c = sum(t["conc"] for t in types); tot_s = sum(t["steel"] for t in types)
    r.add("Totals (per storey height " + fmt(height) + " m)", "", f"{sum(int(c[3]) for c in cols)} columns grouped into {len(types)} types",
          f"concrete {fmt(tot_c, 2)} m³ , longitudinal steel ≈ {fmt(tot_s, 0)} kg")
    r.outputs = dict(types=types, designs=designs, height=height)
    return r


def to_xlsx(r) -> bytes:
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "Column schedule"
    ws.append(["Type", "b (mm)", "h (mm)", "Bars", "Ties", "ρ (%)", "Columns", "Count", "Pu,max (kN)", "Concrete (m³)", "Steel (kg)"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F6FEB"); c.alignment = Alignment(horizontal="center")
    for t in r.outputs["types"]:
        ws.append([t["type"], t["b"], t["h"], f"{t['n']}Ø{t['db']}", f"Ø{t['tie']}@{t['s']}", round(t["rho"] * 100, 2), ", ".join(t["members"]),
                   t["count"], round(t["Pmax"], 1), round(t["conc"], 2), round(t["steel"], 0)])
    ws2 = wb.create_sheet("Columns")
    ws2.append(["Column", "Pu (kN)", "Mu (kN·m)", "Type"])
    tmap = {m: t["type"] for t in r.outputs["types"] for m in t["members"]}
    for name, Pu, Mu, cnt, o in r.outputs["designs"]:
        ws2.append([name, Pu, Mu, tmap[name]])
    for w, col in zip([7, 8, 8, 10, 11, 7, 40, 7, 11, 13, 10], "ABCDEFGHIJK"):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
